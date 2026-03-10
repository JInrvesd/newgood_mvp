from docx import Document
import base64
import io
import re


MAX_EXTRACTED_IMAGES = 20


def _extract_all_images_from_docx(doc) -> list:
    """DOCX에서 모든 이미지(JPEG/PNG)를 base64 data URL 리스트로 추출. (최대 20개)"""
    images = []
    for rel in doc.part.rels.values():
        if not hasattr(rel, "target_part"):
            continue
        try:
            content_type = rel.target_part.content_type
            if any(t in content_type for t in ["jpeg", "png", "jpg"]):
                image_bytes = rel.target_part.blob
                ext = "jpeg" if "jpeg" in content_type or "jpg" in content_type else "png"
                b64 = base64.b64encode(image_bytes).decode("utf-8")
                images.append(f"data:image/{ext};base64,{b64}")
                if len(images) >= MAX_EXTRACTED_IMAGES:
                    break
        except Exception:
            continue
    return images


def parse_docx_to_form_data(file_bytes: bytes) -> dict:
    """
    DOCX 파일을 읽어서 form_data 구조로 파싱.
    Best-effort 방식으로 텍스트 추출 후 _raw_text에 보존.
    LLM이 이후 정제.
    """
    doc = Document(io.BytesIO(file_bytes))
    all_images = _extract_all_images_from_docx(doc)
    photo_data = all_images[0] if all_images else ""

    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    full_text.append(cell.text.strip())

    raw_text = "\n".join(full_text)

    return {
        "personal": {
            "name": _extract_name(raw_text),
            "birthDate": "",
            "phone": _extract_phone(raw_text),
            "email": _extract_email(raw_text),
            "address": "",
            "linkedinUrl": _extract_linkedin(raw_text),
            "portfolioUrl": "",
            "photoData": photo_data,
        },
        "education": [],
        "experience": [],
        "competency": [],
        "others": {
            "certificates": [],
            "languages": [],
            "activities": [],
        },
        "careerDetails": [],
        "coverLetter": {
            "mode": "structured",
            "freeText": "",
            "growth": "",
            "personality": "",
            "motivation": "",
            "aspiration": "",
        },
        "_raw_text": raw_text,
        "_extracted_images": all_images,
    }


def _extract_name(text: str) -> str:
    """한국어 이름(2~5글자 한글) 추출 시도. '이름:' 접두사 패턴도 지원."""
    for line in text.splitlines():
        line = line.strip()
        # "이름: 홍길동" 패턴
        m = re.match(r'이름\s*[:：]\s*([가-힣]{2,5})', line)
        if m:
            return m.group(1)
        # 단독 이름 (2~5자 한글)
        if re.fullmatch(r'[가-힣]{2,5}', line):
            return line
    return ""


def _extract_email(text: str) -> str:
    pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b'
    match = re.search(pattern, text)
    return match.group(0) if match else ""


def _extract_phone(text: str) -> str:
    pattern = r'(?:010|011|016|017|018|019)[-.\s]?\d{3,4}[-.\s]?\d{4}'
    match = re.search(pattern, text)
    return match.group(0) if match else ""


def _extract_linkedin(text: str) -> str:
    pattern = r'linkedin\.com/in/[\w-]+'
    match = re.search(pattern, text, re.IGNORECASE)
    return f"https://{match.group(0)}" if match else ""

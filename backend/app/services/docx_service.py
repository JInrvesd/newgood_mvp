"""
python-docx 기반 이력서 DOCX 생성 서비스.
form_data(7단계 폼)와 result_data(AI 분석 결과)를 받아 .docx 파일을 생성합니다.
"""

import base64
import io
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ---------------------------------------------------------------------------
# 헬퍼 함수
# ---------------------------------------------------------------------------

def _set_cell_bg(cell, hex_color: str):
    """표 셀 배경색 설정 (hex: 예: '2563EB')"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_cell_width(cell, width_cm: float):
    """셀 너비 설정 (cm 단위 → twips 변환)"""
    twips = int(width_cm * 567.17)
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for existing in tcPr.findall(qn("w:tcW")):
        tcPr.remove(existing)
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(twips))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)


def _remove_table_borders(table):
    """테이블 외곽/내부 테두리 제거"""
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement("w:tblBorders")
    for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "none")
        tblBorders.append(border)
    tblPr.append(tblBorders)


def _add_photo_cell_border(cell):
    """사진 셀에 라이트 블루 테두리 추가"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ["top", "left", "bottom", "right"]:
        border = OxmlElement(f"w:{side}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "BFDBFE")
        tcBorders.append(border)
    tcPr.append(tcBorders)


def _add_heading(doc: Document, text: str, level: int = 1):
    """섹션 제목 단락 추가"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    if level == 1:
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)  # blue-600
    else:
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
    # 하단 구분선
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFDBFE")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def _add_label_value(doc: Document, label: str, value: str):
    """레이블: 값 형태의 단락"""
    if not value:
        return
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    label_run = p.add_run(f"{label}: ")
    label_run.bold = True
    label_run.font.size = Pt(10)
    label_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    val_run = p.add_run(value)
    val_run.font.size = Pt(10)


def _add_body(doc: Document, text: str):
    """일반 본문 단락"""
    if not text:
        return
    p = doc.add_paragraph(text)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    for run in p.runs:
        run.font.size = Pt(10)


# ---------------------------------------------------------------------------
# 섹션별 렌더링
# ---------------------------------------------------------------------------

def _render_personal(doc: Document, personal: dict):
    _add_heading(doc, "기본 정보")
    name = personal.get("name") or personal.get("이름", "")
    photo_data = personal.get("photoData", "") or ""

    fields = [
        ("생년월일", personal.get("birthDate") or personal.get("birth_date")),
        ("연락처", personal.get("phone")),
        ("이메일", personal.get("email")),
        ("주소", personal.get("address")),
        ("LinkedIn", personal.get("linkedinUrl") or personal.get("linkedin")),
        ("포트폴리오", personal.get("portfolioUrl") or personal.get("portfolio")),
    ]

    if photo_data:
        # 2열 테이블: 좌=개인정보(11.5cm), 우=사진(4.5cm)
        table = doc.add_table(rows=1, cols=2)
        _remove_table_borders(table)

        left_cell = table.rows[0].cells[0]
        right_cell = table.rows[0].cells[1]
        _set_cell_width(left_cell, 11.5)
        _set_cell_width(right_cell, 4.5)

        # 좌측 셀: 이름
        if name:
            p = left_cell.paragraphs[0]
            run = p.add_run(name)
            run.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
        for label, val in fields:
            if not val:
                continue
            p = left_cell.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            lr = p.add_run(f"{label}: ")
            lr.bold = True
            lr.font.size = Pt(10)
            lr.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            vr = p.add_run(val)
            vr.font.size = Pt(10)

        # 우측 셀: 사진 박스
        _add_photo_cell_border(right_cell)
        _set_cell_bg(right_cell, "EFF6FF")
        p = right_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            data = photo_data
            if data.startswith("data:"):
                data = data.split(",", 1)[1]
            image_bytes = base64.b64decode(data)
            run = p.add_run()
            run.add_picture(io.BytesIO(image_bytes), width=Cm(3.5))
        except Exception as e:
            print(f"[DocxService] 사진 렌더링 실패: {e}")
            run = p.add_run("사 진")
            run.font.color.rgb = RGBColor(0x93, 0xC5, 0xFD)
            run.font.size = Pt(11)
    else:
        # 사진 없음 - 기존 방식
        if name:
            p = doc.add_paragraph()
            run = p.add_run(name)
            run.bold = True
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
            p.paragraph_format.space_after = Pt(4)
        for label, val in fields:
            if val:
                _add_label_value(doc, label, val)


def _render_education(doc: Document, education: list):
    if not education:
        return
    _add_heading(doc, "학력")
    for edu in education:
        school = edu.get("schoolName") or edu.get("school", "")
        major = edu.get("major", "")
        start = edu.get("enrollDate") or edu.get("start_date", "")
        end = edu.get("graduateDate") or edu.get("end_date", "졸업")
        degree = edu.get("degree", "")
        gpa = edu.get("gpa", "")

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        run = p.add_run(school)
        run.bold = True
        run.font.size = Pt(11)

        parts = []
        if major:
            parts.append(major)
        if degree:
            parts.append(degree)
        if parts:
            _add_body(doc, " · ".join(parts))

        period = f"{start} ~ {end}" if start else (end or "")
        if period:
            _add_label_value(doc, "기간", period)
        if gpa:
            _add_label_value(doc, "학점", gpa)


def _calc_duration(start: str, end: str, is_current: bool) -> str:
    """경력 기간을 계산하여 'X년 Y개월' 형식으로 반환"""
    if not start:
        return ""
    try:
        from datetime import date
        s_parts = start.split("-")
        s_year, s_month = int(s_parts[0]), int(s_parts[1]) if len(s_parts) > 1 else 1

        if is_current:
            today = date.today()
            e_year, e_month = today.year, today.month
        elif end:
            e_parts = end.split("-")
            e_year, e_month = int(e_parts[0]), int(e_parts[1]) if len(e_parts) > 1 else 1
        else:
            return ""

        months = (e_year - s_year) * 12 + (e_month - s_month)
        if months < 0:
            return ""
        years = months // 12
        remain = months % 12
        if years == 0 and remain == 0:
            return "1개월 미만"
        if years == 0:
            return f"{remain}개월"
        if remain == 0:
            return f"{years}년"
        return f"{years}년 {remain}개월"
    except Exception:
        return ""


def _render_experience(doc: Document, experience: list, career_details: list = None):
    if not experience:
        return
    _add_heading(doc, "경력")
    details_map = {}
    if career_details:
        for d in career_details:
            key = d.get("companyName", "")
            details_map[key] = d

    for exp in experience:
        company = exp.get("companyName") or exp.get("company", "")
        position = exp.get("position", "")
        start = exp.get("startDate") or exp.get("start_date", "")
        end = exp.get("endDate") or exp.get("end_date", "")
        is_current = exp.get("isCurrent") or exp.get("is_current", False)
        description = exp.get("description", "")

        # 기간 계산
        duration = _calc_duration(start, end, is_current)

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        run = p.add_run(company)
        run.bold = True
        run.font.size = Pt(11)
        # 회사명 옆에 근무 기간 표시
        if duration:
            dur_run = p.add_run(f"  ({duration})")
            dur_run.font.size = Pt(9)
            dur_run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

        if position:
            _add_label_value(doc, "직위", position)

        end_label = "재직 중" if is_current else (end or "")
        period = f"{start} ~ {end_label}" if start else end_label
        if period:
            _add_label_value(doc, "기간", period)

        if description:
            _add_body(doc, description)

        # 경력 상세 (career_details)
        detail = details_map.get(company, {})
        if detail.get("achievements"):
            _add_label_value(doc, "주요 성과", detail["achievements"])
        if detail.get("projectName"):
            _add_label_value(doc, "프로젝트", detail["projectName"])
        if detail.get("skills"):
            skills = detail["skills"]
            if isinstance(skills, list):
                skills = ", ".join(skills)
            _add_label_value(doc, "사용 기술", skills)


def _render_competency(doc: Document, competency: list):
    if not competency:
        return
    _add_heading(doc, "역량 / 스킬")
    for item in competency:
        name = item.get("name", "")
        level = item.get("level", "")
        desc = item.get("description", "")
        if not name:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        run = p.add_run(f"• {name}")
        run.bold = True
        run.font.size = Pt(10)
        if level:
            run2 = p.add_run(f"  [{level}]")
            run2.font.size = Pt(9)
            run2.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        if desc:
            _add_body(doc, f"  {desc}")


def _render_others(doc: Document, others: dict):
    if not others:
        return

    certs = others.get("certificates", [])
    langs = others.get("languages", [])
    acts = others.get("activities", [])

    valid_certs = [c for c in certs if c.get("name")]
    valid_langs = [l for l in langs if l.get("language")]
    valid_acts = [a for a in acts if a.get("name") or a.get("title")]

    if not (valid_certs or valid_langs or valid_acts):
        return

    _add_heading(doc, "자격증 / 어학 / 활동")

    if valid_certs:
        p = doc.add_paragraph()
        run = p.add_run("자격증")
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        for c in valid_certs:
            parts = [c["name"]]
            if c.get("issuer"):
                parts.append(c["issuer"])
            if c.get("date"):
                parts.append(c["date"])
            _add_body(doc, "  • " + "  |  ".join(parts))

    if valid_langs:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        run = p.add_run("어학")
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        for l in valid_langs:
            parts = [l["language"]]
            if l.get("level"):
                parts.append(l["level"])
            if l.get("score"):
                parts.append(l["score"])
            _add_body(doc, "  • " + "  |  ".join(parts))

    if valid_acts:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        run = p.add_run("기타 활동")
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        for a in valid_acts:
            title = a.get("name") or a.get("title", "")
            date = a.get("date", "")
            desc = a.get("description", "")
            line = f"  • {title}"
            if date:
                line += f"  ({date})"
            _add_body(doc, line)
            if desc:
                _add_body(doc, f"    {desc}")


def _render_cover_letter(doc: Document, cover_letter: dict):
    if not cover_letter:
        return

    mode = cover_letter.get("mode", "structured")

    if mode == "free":
        free_text = cover_letter.get("freeText") or cover_letter.get("free_text", "")
        if not free_text:
            return
        _add_heading(doc, "자기소개서")
        _add_body(doc, free_text)
        return

    # structured mode
    sections = [
        ("성장과정", cover_letter.get("growth", "")),
        ("성격 / 장단점", cover_letter.get("personality", "")),
        ("지원동기", cover_letter.get("motivation", "")),
        ("입사 후 포부", cover_letter.get("aspiration", "")),
    ]
    has_content = any(v for _, v in sections)
    if not has_content:
        return

    _add_heading(doc, "자기소개서")
    for label, text in sections:
        if not text:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        run = p.add_run(label)
        run.bold = True
        run.font.size = Pt(11)
        _add_body(doc, text)


# ---------------------------------------------------------------------------
# 메인 생성 함수
# ---------------------------------------------------------------------------

async def generate_docx(form_data: dict, result_data: dict) -> bytes:
    """
    form_data와 result_data를 기반으로 DOCX 파일을 생성합니다.
    python-docx를 직접 사용합니다.
    """
    doc = Document()

    # 페이지 여백 설정
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # 기본 폰트 설정
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10)

    # 섹션 렌더링
    personal = form_data.get("personal", {})
    education = form_data.get("education", [])
    experience = form_data.get("experience", [])
    career_details = form_data.get("careerDetails", [])
    competency = form_data.get("competency", [])
    others = form_data.get("others", {})
    cover_letter = form_data.get("coverLetter") or form_data.get("cover_letter", {})

    _render_personal(doc, personal)
    _render_education(doc, education)
    _render_experience(doc, experience, career_details)
    _render_competency(doc, competency)
    _render_others(doc, others)
    _render_cover_letter(doc, cover_letter)

    # AI 개선 제안이 있으면 부록으로 추가
    phase1 = result_data.get("phase1", {})
    structured = phase1.get("structured_resume", {})
    highlights = structured.get("experience_highlights", [])
    if highlights:
        _add_heading(doc, "[AI 개선 제안] 경력 기술 개선안")
        for h in highlights:
            company = h.get("company", "")
            improved = h.get("improved_description", "")
            if company and improved:
                p = doc.add_paragraph()
                run = p.add_run(company)
                run.bold = True
                run.font.size = Pt(10)
                _add_body(doc, improved)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
